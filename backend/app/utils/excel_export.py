"""One way to turn admin lists into an Excel download, so every export looks and behaves the same."""
import io
from datetime import datetime, timezone
from typing import Iterable, Optional, Sequence

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.controllers.cookie_consent import csv_safe
from app.utils.timestamps import IST

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def ist_parts(value: Optional[datetime]) -> tuple[str, str]:
    """('29-10-2026', '03:42 PM') in Indian time, or two empty strings."""
    if not value:
        return "", ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    moment = value.astimezone(IST)
    return moment.strftime("%d-%m-%Y"), moment.strftime("%I:%M %p")


def ist_text(value: Optional[datetime]) -> str:
    date, time = ist_parts(value)
    return f"{date} {time}".strip()


def build_xlsx(sheet_title: str, headers: Sequence[str], rows: Iterable[Sequence], widths: Optional[Sequence[int]] = None) -> tuple[bytes, int]:
    """Excel file (bytes) and the number of data rows. Text that starts with = + - @ is neutralised so Excel never runs it."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append(list(headers))
    navy = PatternFill("solid", fgColor="0F2340")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill, cell.alignment = Font(bold=True, color="FFFFFF"), navy, Alignment(vertical="center", wrap_text=True)
    count = 0
    for row in rows:
        ws.append([csv_safe(v) if isinstance(v, str) else v for v in row])
        count += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(count + 1, 2)}"
    for i, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths[i - 1] if widths and i <= len(widths) else max(12, min(40, len(header) + 4)))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), count


def xlsx_response(body: bytes, count: int, filename: str) -> Response:
    stamp = datetime.now(IST).strftime("%d-%m-%Y")
    name = filename.replace("{date}", stamp)
    return Response(content=body, media_type=XLSX_TYPE, headers={
        "Content-Disposition": f'attachment; filename="{name}"', "Cache-Control": "no-store",
        "X-Content-Type-Options": "nosniff", "X-Application-Count": str(count),
    })
