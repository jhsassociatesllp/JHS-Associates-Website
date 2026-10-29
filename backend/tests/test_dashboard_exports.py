"""Run: python tests/test_dashboard_exports.py (no database needed).

Live dashboard numbers (role-scoped) and the Excel downloads for Contacts, Users and Client feedback.
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from datetime import datetime, timedelta, timezone

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import deps
from app.auth.security import create_access_token
from app.main import app
from app.routes import dashboard as dash
from app.routes import exports as exp

IST = timezone(timedelta(hours=5, minutes=30))


def at(y, m, d, h, mi):
    return datetime(y, m, d, h, mi, tzinfo=IST).astimezone(timezone.utc)


# ── fake database for the exports ──
DATA = {
    "contact_us": [
        {"name": "Asha Rao", "email": "asha@x.com", "phone": "9876543210", "company": "Acme", "service": "Consulting", "message": "Hello there, need help", "created_at": at(2026, 10, 22, 23, 50)},
        {"name": "=HYPERLINK(\"x\")", "email": "evil@x.com", "message": "formula test message", "created_at": at(2026, 10, 24, 9, 0)},
    ],
    "users": [
        {"name": "Hitman", "email": "h@x.com", "auth_provider": "google", "is_verified": True, "status": "active", "created_at": at(2026, 10, 23, 10, 0), "last_login": at(2026, 10, 24, 11, 0)},
        {"first_name": "Dev", "last_name": "Patel", "email": "d@x.com", "auth_provider": "password", "created_at": at(2026, 10, 25, 10, 0)},
    ],
    "feedback": [
        {"client_name": "Aga Khan Rural Support", "name": "Ravi", "designation": "CFO", "nature_of_assignment": "IA", "overall": 4, "would_refer": "Yes",
         "references": [{"name_designation": "Sunil, MD", "company_name": "Zenda", "phone": "99", "email": "s@z.com"}], "created_at": at(2026, 10, 16, 18, 10)},
    ],
}


class Cur:
    def __init__(self, rows):
        self.rows = rows

    def __aiter__(self):
        self._i = iter(self.rows)
        return self

    async def __anext__(self):
        try:
            return next(self._i)
        except StopIteration:
            raise StopAsyncIteration


class Col:
    def __init__(self, name):
        self.name = name

    def find(self, q=None):
        q = q or {}
        rows = []
        for d in DATA.get(self.name, []):
            c = q.get("created_at")
            if c:
                t = d["created_at"].astimezone(timezone.utc)
                if "$gte" in c and t < c["$gte"]:
                    continue
                if "$lt" in c and t >= c["$lt"]:
                    continue
            rows.append(dict(d))
        return Cur(rows)


class DB(dict):
    def __getitem__(self, k):
        return Col(k)


exp.get_database = lambda: DB()
deps.get_database = lambda: DB()


async def fake_admin(email):
    role = {"boss@x.com": "super_admin", "hr@x.com": "hr_admin", "editor@x.com": "admin"}[email]
    return {"_id": "a1", "email": email, "password": "x", "role": role, "name": "Team"}


deps.find_admin_by_email = fake_admin
c = TestClient(app, raise_server_exceptions=False)


def tok(email, role):
    return {"Authorization": "Bearer " + create_access_token({"sub": email, "role": role, "type": "admin"}, expires_delta=timedelta(hours=1))}


boss, hr, editor = tok("boss@x.com", "super_admin"), tok("hr@x.com", "hr_admin"), tok("editor@x.com", "admin")
ok = True


def check(name, cond, extra=""):
    global ok
    ok &= bool(cond)
    print(("PASS " if cond else "FAIL ") + name, extra)


def sheet(resp):
    return [[cell for cell in row] for row in load_workbook(io.BytesIO(resp.content)).active.iter_rows(values_only=True)]


# ── Excel downloads ──
for path in ("contacts", "users", "feedback"):
    check(f"{path} export: no login -> 401", c.get(f"/api/export/{path}").status_code == 401)
    check(f"{path} export: content admin -> 403", c.get(f"/api/export/{path}", headers=editor).status_code == 403)

r = c.get("/api/export/contacts", headers=hr)
rows = sheet(r)
check("contacts export: Excel with every contact, newest first", r.status_code == 200 and "spreadsheetml" in r.headers["content-type"] and len(rows) == 3 and rows[1][3] == "=HYPERLINK(\"x\")".join(["'", ""]) or rows[1][3].startswith("'="), [x[3] for x in rows[1:]])
check("contacts export: Indian date and time, formula neutralised", rows[2][1] == "22-10-2026" and rows[2][2] == "11:50 PM" and str(rows[1][3]).startswith("'="), rows[2][:3])
r = c.get("/api/export/contacts?date_from=2026-10-23&date_to=2026-10-24", headers=boss)
check("contacts export: date filter (Indian days)", r.headers["x-application-count"] == "1")
check("contacts export: bad date -> 422", c.get("/api/export/contacts?date_from=hello", headers=boss).status_code == 422)

rows = sheet(c.get("/api/export/users", headers=hr))
check("users export: name, sign-in method, IST times", len(rows) == 3 and rows[1][1] == "Dev Patel" and rows[1][3] == "Email & password" and rows[2][3] == "Google" and rows[2][6] == "23-10-2026 10:00 AM", rows[2])
rows = sheet(c.get("/api/export/feedback", headers=hr))
check("feedback export: client, submitter, rating, references", rows[1][3] == "Aga Khan Rural Support" and rows[1][4] == "Ravi" and "Sunil, MD" in str(rows[1][-1]) and 4 in rows[1], rows[1][:5])

# ── Dashboard numbers, scoped by role ──
async def fake_facet(collection, w, attention):
    base = {"total": [{"n": 10}], "today": [{"n": 2}], "week": [{"n": 5}], "month": [{"n": 7}],
            "by_day": [{"_id": w["days"][-1], "n": 2}], "by_month": [{"_id": w["months"][-1], "n": 7}]}
    if attention:
        base["attention"] = [{"n": 3}]
    return base


async def fake_recent(kind, collection, limit=6):
    return [{"type": kind, "title": f"{kind} item", "subtitle": "", "at": datetime(2026, 10, 29, 6, 0, tzinfo=timezone.utc)}]


dash._facet, dash._recent = fake_facet, fake_recent
check("dashboard: no login -> 401", c.get("/api/dashboard/stats").status_code == 401)

d = c.get("/api/dashboard/stats", headers=boss).json()
check("Super Admin gets content and submissions", d["content"] and d["submissions"] and d["role"] == "super_admin")
check("content has Regulatory, Reg-Watch, White Papers (not just blogs)", {"regulatory", "reg_watch", "white_papers", "articles", "knowledge"} <= set(d["content"]), list(d["content"]))
check("submission counts: total, today, week, month, waiting for us", d["submissions"]["proposals"]["total"] == 10 and d["submissions"]["proposals"]["attention"] == 3 and "attention" not in d["submissions"]["contacts"], d["submissions"]["proposals"])
check("series: 7 days, 6 months and 12 months", len(d["submission_series"]["week"]) == 7 and len(d["content_series"]["month"]) == 6 and len(d["content_series"]["year"]) == 12)
check("recent feed is newest-first with admin links and UTC times", d["recent"] and d["recent"][0]["at"].endswith("Z") and d["recent"][0]["path"].startswith("/"), d["recent"][0])
d = c.get("/api/dashboard/stats", headers=hr).json()
check("HR Admin: submissions only, no content", d["submissions"] and d["content"] is None and d["content_series"] is None)
d = c.get("/api/dashboard/stats", headers=editor).json()
check("content admin: content only, no submissions or form data", d["content"] and d["submissions"] is None and all(i["type"] in dash.CONTENT for i in d["recent"]))


async def broken(collection, w, attention):
    raise RuntimeError("database hiccup")


dash._facet = broken
d = c.get("/api/dashboard/stats", headers=boss)
check("one failing collection never blanks the dashboard", d.status_code == 200 and d.json()["content"]["articles"]["total"] == 0)

# window maths: Indian days / months
w = dash._windows(datetime(2026, 10, 28, 19, 0, tzinfo=timezone.utc))   # = 29 Oct, 00:30 IST
check("window: just after Indian midnight counts as the new day", w["days"][-1] == "2026-10-29" and w["months"][-1] == "2026-10" and w["months"][0] == "2025-11", (w["days"][-1], w["months"][0]))
print("ALL PASS" if ok else "SOME FAILED")
