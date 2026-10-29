"""Run: python tests/test_proposal_admin.py (no database needed).

Admin tracking of Request-for-Proposal submissions: list, status + note, Excel export, admin-only.
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from datetime import datetime, timedelta, timezone

from bson import ObjectId
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import deps
from app.auth.security import create_access_token
from app.controllers import proposal as pc
from app.main import app

IST = timezone(timedelta(hours=5, minutes=30))


def at(y, m, d, h, mi):
    return datetime(y, m, d, h, mi, tzinfo=IST).astimezone(timezone.utc)


def doc(first, created, **extra):
    return {"_id": ObjectId(), "first_name": first, "last_name": "Test", "email": f"{first.lower()}@x.com", "phone": "9876543210",
            "inquiry_reason": "Request for services", "subject": "Audit help", "message": "Need an audit proposal",
            "user_id": "u1", "created_at": created, **extra}


DOCS = [doc("Asha", at(2026, 10, 22, 23, 50)),                      # no status yet (older request)
        doc("Bilal", at(2026, 10, 23, 0, 5), status="in_review"),
        doc("Chitra", at(2026, 10, 23, 15, 30), status="new"),
        {**doc("Esha", at(2026, 10, 24, 9, 0), status="closed"), "first_name": "=HYPERLINK(\"x\")"}]


class Cur:
    def __init__(self, rows):
        self.rows = rows

    def sort(self, *a, **k):
        self.rows = sorted(self.rows, key=lambda d: d["created_at"], reverse=True)
        return self

    def __aiter__(self):
        self._i = iter(self.rows)
        return self

    async def __anext__(self):
        try:
            return next(self._i)
        except StopIteration:
            raise StopAsyncIteration


def matches(d, q):
    c = q.get("created_at")
    if c:
        t = d["created_at"].astimezone(timezone.utc)
        if "$gte" in c and t < c["$gte"]:
            return False
        if "$lt" in c and t >= c["$lt"]:
            return False
    if "status" in q and d.get("status") != q["status"]:
        return False
    if "$or" in q and not any(("status" in o and o["status"] == "new" and d.get("status") == "new") or ("status" in o and isinstance(o["status"], dict) and "status" not in d) for o in q["$or"]):
        return False
    return True


class Col:
    def find(self, q=None):
        return Cur([dict(d) for d in DOCS if matches(d, q or {})])

    async def find_one(self, q):
        for d in DOCS:
            if d["_id"] == q["_id"]:
                return dict(d)

    async def update_one(self, q, upd):
        class R:
            matched_count = 0
        r = R()
        for d in DOCS:
            if d["_id"] == q["_id"]:
                d.update(upd["$set"])
                r.matched_count = 1
        return r


class DB(dict):
    def __getitem__(self, k):
        return Col()


pc.get_database = lambda: DB()
deps.get_database = lambda: DB()


async def fake_admin(email):
    role = "hr_admin" if email == "hr@x.com" else "admin"
    return {"_id": "a1", "email": email, "password": "x", "role": role, "name": "Team"}


deps.find_admin_by_email = fake_admin
c = TestClient(app, raise_server_exceptions=False)


def tok(email, role):
    return {"Authorization": "Bearer " + create_access_token({"sub": email, "role": role, "type": "admin"}, expires_delta=timedelta(hours=1))}


hr, content_admin = tok("hr@x.com", "hr_admin"), tok("editor@x.com", "admin")
ok = True


def check(name, cond, extra=""):
    global ok
    ok &= bool(cond)
    print(("PASS " if cond else "FAIL ") + name, extra)


U = "/api/proposal"
check("list: no login -> 401", c.get(U + "/").status_code == 401)
check("list: visitor token -> 401", c.get(U + "/", headers={"Authorization": "Bearer " + create_access_token({"sub": "v@x.com", "uid": "u", "type": "site_user"}, expires_delta=timedelta(hours=1))}).status_code == 401)
check("list: content admin (not HR) -> 403", c.get(U + "/", headers=content_admin).status_code == 403)
check("status change: no login -> 401", c.patch(f"{U}/{DOCS[0]['_id']}/status", json={"status": "closed"}).status_code == 401)
check("status change: content admin -> 403", c.patch(f"{U}/{DOCS[0]['_id']}/status", json={"status": "closed"}, headers=content_admin).status_code == 403)
check("export: no login -> 401", c.get(U + "/export").status_code == 401)
check("export: content admin -> 403", c.get(U + "/export", headers=content_admin).status_code == 403)

r = c.get(U + "/", headers=hr)
rows = r.json()
check("HR admin sees every request, older ones default to 'new'", r.status_code == 200 and len(rows) == 4 and [x for x in rows if x["first_name"] == "Asha"][0]["status"] == "new", r.status_code)

aid = str(DOCS[0]["_id"])
r = c.patch(f"{U}/{aid}/status", json={"status": "proposal_sent", "note": "  Sent the proposal on 25 Oct  "}, headers=hr)
check("status + note saved, with who and when", r.status_code == 200 and r.json()["status"] == "proposal_sent" and r.json()["admin_note"] == "Sent the proposal on 25 Oct" and r.json()["handled_by"] == "hr@x.com" and r.json()["updated_at"], r.status_code)
check("unknown status -> 422", c.patch(f"{U}/{aid}/status", json={"status": "hacked"}, headers=hr).status_code == 422)
check("markup in the note -> 422", c.patch(f"{U}/{aid}/status", json={"status": "closed", "note": "<script>x</script>"}, headers=hr).status_code == 422)
check("note too long -> 422", c.patch(f"{U}/{aid}/status", json={"status": "closed", "note": "a" * 1001}, headers=hr).status_code == 422)
check("bad id -> 400", c.patch(f"{U}/not-an-id/status", json={"status": "closed"}, headers=hr).status_code == 400)
check("missing request -> 404", c.patch(f"{U}/{ObjectId()}/status", json={"status": "closed"}, headers=hr).status_code == 404)

r = c.get(U + "/export?date_from=2026-10-23&date_to=2026-10-23", headers=hr)
wb = load_workbook(io.BytesIO(r.content))
sheet = [[x for x in row] for row in wb.active.iter_rows(values_only=True)]
check("export: Excel, only that Indian day (Bilal + Chitra)", r.status_code == 200 and "spreadsheetml" in r.headers["content-type"] and r.headers["x-application-count"] == "2" and [s[3] for s in sheet[1:]] == ["Bilal", "Chitra"], [s[3] for s in sheet[1:]])
check("export: IST date and readable status", sheet[1][1] == "23-10-2026" and sheet[1][2] == "12:05 AM" and sheet[1][10] == "In review", sheet[1][:11])
r = c.get(U + "/export", headers=hr)
names = [x[3] for x in [[c_ for c_ in row] for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True)][1:]]
check("export: formula in a name neutralised", any(str(n).startswith("'=") for n in names), names)
r = c.get(U + "/export?status=closed", headers=hr)
check("export: status filter", r.headers["x-application-count"] == "1")
check("export: bad date -> 422", c.get(U + "/export?date_from=hello", headers=hr).status_code == 422)
check("export: unknown status -> 422", c.get(U + "/export?status=zzz", headers=hr).status_code == 422)
print("ALL PASS" if ok else "SOME FAILED")
