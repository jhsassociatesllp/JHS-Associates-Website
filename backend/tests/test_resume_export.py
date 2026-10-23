"""Run: python tests/test_resume_export.py (no database needed).

Daily résumé report: date filter (Indian days), ZIP with every résumé + report, admin-only.
"""
import io
import os
import sys
import zipfile

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from datetime import datetime, timedelta, timezone

from bson import ObjectId
from fastapi.testclient import TestClient

from app.auth import deps
from app.auth.security import create_access_token
from app.controllers import career_export as ce
from app.main import app

IST = timezone(timedelta(hours=5, minutes=30))


def at(y, m, d, h, mi):
    """An Indian clock time, stored the way the app stores it (UTC)."""
    return datetime(y, m, d, h, mi, tzinfo=IST).astimezone(timezone.utc)


FILES = {}
APPS = []


def add(name, created, job="Senior Auditor", job_id="j1", with_file=True, extra=None):
    fid = ObjectId()
    if with_file:
        FILES[str(fid)] = b"%PDF-1.4 resume of " + name.encode()
    APPS.append({"_id": ObjectId(), "full_name": name, "email": f"{name.split()[0].lower()}@x.com", "phone": "9876543210", "job_id": job_id,
                 "job_title": job, "created_at": created, "resume_file_id": str(fid) if with_file else None, "status": "new",
                 "ats_score": 72, "ats_status": "matched", "highest_qualification": "Graduate", "profile": "Fresher", "how_heard": "LinkedIn", **(extra or {})})


add("Asha Rao", at(2026, 10, 22, 23, 50))            # 22 Oct, just before midnight IST
add("Bilal Khan", at(2026, 10, 23, 0, 5))            # 23 Oct, just after midnight IST
add("Chitra Nair", at(2026, 10, 23, 15, 30))
add("=HYPERLINK(\"x\")", at(2026, 10, 23, 18, 0), job="General", job_id=None, with_file=False)
add("Esha Mehta", at(2026, 10, 24, 9, 0))


class Col:
    def find(self, q):
        class Cur:
            def __init__(s, rows): s.rows = rows
            def __aiter__(s): s._i = iter(s.rows); return s
            async def __anext__(s):
                try: return next(s._i)
                except StopIteration: raise StopAsyncIteration
        rows = []
        for d in APPS:
            c = q.get("created_at")
            if c:
                t = d["created_at"].astimezone(timezone.utc)
                if "$gte" in c and t < c["$gte"]: continue
                if "$lt" in c and t >= c["$lt"]: continue
            if "job_id" in q and d["job_id"] != q["job_id"]: continue
            if "$or" in q and d["job_id"]: continue
            rows.append(dict(d))
        return Cur(rows)


class DB(dict):
    def __getitem__(self, k): return Col()


class FakeBucket:
    def __init__(self, db, bucket_name=None): pass
    async def open_download_stream(self, oid):
        data = FILES.get(str(oid))
        if data is None: raise FileNotFoundError("missing")
        class S:
            async def read(s): return data
        return S()


ce.get_database = lambda: DB()
ce.AsyncIOMotorGridFSBucket = FakeBucket
deps.get_database = lambda: DB()


async def fake_admin(email):
    return {"_id": "a1", "email": email, "password": "x", "role": "hr_admin", "name": "HR"}


deps.find_admin_by_email = fake_admin
c = TestClient(app, raise_server_exceptions=False)
hr = {"Authorization": "Bearer " + create_access_token({"sub": "hr@x.com", "role": "hr_admin", "type": "admin"}, expires_delta=timedelta(hours=1))}
ok = True


def check(name, cond, extra=""):
    global ok
    ok &= bool(cond)
    print(("PASS " if cond else "FAIL ") + name, extra)


U = "/api/careers/admin/applications/export"
check("no login -> 401", c.get(U + "?date_from=2026-10-23&date_to=2026-10-23").status_code == 401)
check("visitor token -> 401", c.get(U, headers={"Authorization": "Bearer " + create_access_token({"sub": "v@x.com", "uid": "u", "type": "site_user"}, expires_delta=timedelta(hours=1))}).status_code == 401)

r = c.get(U + "?date_from=2026-10-23&date_to=2026-10-23", headers=hr)
check("one day -> ZIP", r.status_code == 200 and r.headers["content-type"] == "application/zip" and "JHS_Resumes_23-10-2026.zip" in r.headers["content-disposition"], r.status_code)
z = zipfile.ZipFile(io.BytesIO(r.content))
names = z.namelist()
root = "JHS_Resumes_23-10-2026/"
pdfs = [n for n in names if n.endswith(".pdf")]
check("only that Indian day (midnight edges handled): Bilal + Chitra, not Asha / Esha", r.headers["x-application-count"] == "3" and len(pdfs) == 2 and any("Bilal" in n for n in pdfs) and any("Chitra" in n for n in pdfs) and not any("Asha" in n or "Esha" in n for n in pdfs), (r.headers["x-application-count"], pdfs))
check("folder has Tracker, read-me", all(root + f in names for f in ("Tracker.xlsx", "READ_ME.txt")))
wb = load_workbook(io.BytesIO(z.read(root + "Tracker.xlsx")))
rows = [[x for x in r] for r in wb["Tracker"].iter_rows(values_only=True)]
check("tracker: header + 3 rows, IST date/time", len(rows) == 4 and rows[1][1] == "23-10-2026" and rows[1][2] == "12:05 AM", rows[1][:4])
check("candidate without a file is listed too", any("(no resume file)" in str(x) for row in rows for x in row))
check("spreadsheet formula in a name is neutralised", any(str(row[3]).startswith("'=") for row in rows[1:]), [row[3] for row in rows[1:]])
summ = [[x for x in r] for r in wb["Summary"].iter_rows(values_only=True)]
flat = [x for r in summ for x in r if x is not None]
check("summary: total, day-wise, by profile", "Total applications: 3" in flat and "23-10-2026" in flat and "Fresher" in flat and "By profile (who applied)" in flat)
check("resume content is intact", z.read([n for n in names if "Bilal" in n][0]).startswith(b"%PDF-1.4 resume of Bilal"))

r2 = c.get(U + "?date_from=2026-10-22&date_to=2026-10-24", headers=hr)
check("three days -> 5 applications, folders per day", r2.headers["x-application-count"] == "5" and {n.split("/")[2] for n in zipfile.ZipFile(io.BytesIO(r2.content)).namelist() if "/Resumes/" in n} == {"22-10-2026", "23-10-2026", "24-10-2026"})
r3 = c.get(U + "?what=report&date_from=2026-10-23&date_to=2026-10-23", headers=hr)
check("report only -> Excel", r3.status_code == 200 and "spreadsheetml" in r3.headers["content-type"] and r3.content[:2] == b"PK" and ".xlsx" in r3.headers["content-disposition"])
r4 = c.get(U + "?date_from=2026-10-23&date_to=2026-10-23&job_id=general", headers=hr)
check("general-applications filter", r4.headers["x-application-count"] == "1")
check("bad date -> 422", c.get(U + "?date_from=hello", headers=hr).status_code == 422)
check("'to' before 'from' -> 422", c.get(U + "?date_from=2026-10-24&date_to=2026-10-22", headers=hr).status_code == 422)
check("bad 'what' -> 422", c.get(U + "?what=exe", headers=hr).status_code == 422)
ce.MAX_RESUMES_PER_DOWNLOAD = 2
check("too many for one download -> 400 with a clear message", c.get(U + "?date_from=2026-10-22&date_to=2026-10-24", headers=hr).status_code == 400)
print("ALL PASS" if ok else "SOME FAILED")
